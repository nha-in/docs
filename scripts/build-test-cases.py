#!/usr/bin/env python3
"""Write the HIE-CM test case pages under Developer resources from NHA's sheets.

NHA publishes each milestone's functional test cases as a spreadsheet. The
sheets are stored untouched under catalogue/hiecm/openapi/.raw/, and this
turns each one into a page with one table per group of cases, the way the
sheet lays them out. A case id is an anchor on its row, so it is searchable,
linkable from a milestone or API page, and present in the page's Markdown copy
for agents. Steps sit behind a Steps toggle in the row, so the table stays
scannable. The pages carry `generated: true`: fix
the sheet or this script, never the page.

Each call a case names links to its endpoint page when the path is exactly
one operation in site/src/data/api-routes.json. A path several operations
share, told apart only by the request body's scope, stays plain code rather
than linking to a guess.

Usage: python3 scripts/build-test-cases.py [--check]
Needs openpyxl.
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "catalogue" / "hiecm" / "openapi" / ".raw"
ROUTES = ROOT / "site" / "src" / "data" / "api-routes.json"
OUT = ROOT / "site" / "docs" / "hiecm" / "v3" / "resources" / "test-cases"

PAGES = {
    "m1": {
        "sheet": RAW / "nha-2026-09-29-m1-tests" / "M1_ABHA_CREATION_AND_VERIFICATION_WITH_APIS_UPDATED_V1_1.xlsx",
        "title": "M1 test cases",
        "label": "M1 Identity",
        "position": 2,
        "milestone": "[M1 Identity](/docs/hiecm/v3/milestones/m1)",
        "covers": "ABHA creation by Aadhaar OTP, Aadhaar biometric, demographic authentication and driving licence or PAN, then verification, profile update, one ABHA per patient record, and sharing a profile by QR code",
        "api": "/docs/hiecm/v3/api/m1",
        "skill": "/docs/hiecm/v3/milestones/m1#build-m1-with-an-ai-coding-assistant",
        "next": "[M2 test cases](/docs/hiecm/v3/resources/test-cases/m2)",
    },
    # M4 comes as two workbooks, HFR with one tab per flow and HPR with one.
    # Each tab is a section of the page, and its heading rows are the groups.
    "m4": {
        "sections": [
            ("HFR", RAW / "nha-2026-10-02-m4-tests" / "HFR_Test_Cases.xlsx"),
            ("HPR", RAW / "nha-2026-10-02-m4-tests" / "HPR_Test_Cases.xlsx"),
        ],
        "title": "M4 test cases",
        "label": "M4 Registry Integration",
        "position": 5,
        "milestone": "[M4 Registry Integration](/docs/hiecm/v3/milestones/m4)",
        "covers": "searching the Health Facility Registry, registering a facility in four steps, updating it and linking bridges to it, then creating a Healthcare Professional ID and registering the professional on the HPR",
        "api": "/docs/hiecm/v3/api/m4",
        "skill": "/docs/hiecm/v3/milestones/m4#build-m4-with-an-ai-coding-assistant",
        "next": "[PHR test cases](/docs/hiecm/v3/resources/test-cases/phr)",
    },
}

# The HFR sheet names its calls by their sandbox swagger operation rather than
# by path. These are the ones that are one M4 operation by name; any other
# stays plain code rather than linking to a guess.
SWAGGER = {
    "v15FacilityBasicInformation": "/v1.5/facility/basic-information",
    "v15FacilityAdditionalInformation": "/v1.5/facility/additional-information",
    "v15FacilityDetailedInformation": "/v1.5/facility/detailed-information",
    "v15SubmitFacilityDetails": "/v1.5/facility/submit-facility",
    "v15FacilityGetLGDStates": "/v1.5/facility/lgd/states",
    "v15FacilityGetLGDDistricts": "/v1.5/facility/lgd/districts",
    "v15FacilityGetLGDSubDistricts": "/v1.5/facility/lgd/subdistricts",
    "v15FacilityGetMasterData": "/v1.5/facility/get-master-data",
    "v15FacilityGetMasterTypes": "/v1.5/facility/get-master-types",
    "v15FacilityGetOwnershipSubtype": "/v1.5/facility/get-owner-subtype",
    "v15FacilityGetSpecialities": "/v1.5/facility/get-specialities",
    "v15FetchFacilityTypes": "/v1.5/facility/fetch-facility-type",
    "v15FetchFacilitySubTypes": "/v1.5/facility/fetch-facility-Sub-type",
    "v1MutipleHRPAddUpdateServices": "/v1/bridges/MutipleHRPAddUpdateServices",
}

# The M4 sheets mark a case Yes or No where M1 writes Mandatory or Optional.
MARKINGS = {"yes": "Mandatory", "no": "Optional"}


def clean(value):
    if value is None:
        return ""
    text = str(value).replace("\r", "").replace("\u2014", ", ")
    lines = [re.sub(r"\s{2,}", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def flat(value):
    return " ".join(clean(value).split("\n"))


def md(text):
    """Escape what CommonMark would read as markup in NHA's prose.

    A leading "1." is escaped too, or a step inside a bullet renders as a
    numbered list nested in it.
    """
    text = re.sub(r"([\\`*\[\]{}<>])", r"\\\1", text)
    return re.sub(r"^(\d+)\.", r"\1\\.", text)


def columns(header):
    found = {}
    for index, name in enumerate(header):
        key = flat(name).lower()
        if not key:
            continue
        if "test case id" in key:
            found["id"] = index
        elif "functionality" in key:
            found["functionality"] = index
        elif key.startswith("function"):
            found["note"] = index
        elif "applicable" in key:
            found["applies"] = index
        elif "mandatory" in key:
            found["marking"] = index
        elif key.startswith("test case"):
            found["test"] = index
        elif "steps" in key:
            found["steps"] = index
        elif "expected" in key:
            found["expected"] = index
        elif "suggestion" in key:
            found["tester"] = index
        elif "v3 api" in key or key == "apis":
            found["apis"] = index
        elif key.startswith("actual name"):
            found["field"] = index
    return found


def load_routes():
    by_path = {}
    for entry in json.loads(ROUTES.read_text()):
        if not entry.get("route", "").startswith("/docs/hiecm/"):
            continue
        bare = entry["path"].split("#")[0]
        by_path.setdefault(bare, []).append(entry)
    return by_path


def call(url, routes, base=""):
    """An API as code, linked when its path is exactly one operation.

    `base` is a prefix every case on the page shares; it is stated once above
    the tables and left off each path, so the APIs column stays narrow.
    """
    parsed = urlparse(url)
    path = parsed.path.rstrip("/") or url
    if path.endswith("swagger-ui.html") and parsed.fragment:
        name = re.sub(r"Using(GET|POST|PUT|PATCH|DELETE)$", "", parsed.fragment.split("/")[-1])
        if name not in SWAGGER:
            return f"`{name}`"
        path = SWAGGER[name]
    shown = path[len(base):] if base and path.startswith(base + "/") else path
    matches = routes.get(path, [])
    if len(matches) == 1:
        return f"[`{matches[0]['method']} {shown}`]({matches[0]['route']})"
    return f"`{shown}`"


def shared_base(groups):
    """The first three path segments most of the page's APIs share.

    Paths outside it, such as a PHR share call on an M1 page, stay in full.
    """
    heads = [
        "/" + "/".join(parts[:3])
        for g in groups for c in g["cases"] for u in c["apis"]
        if len(parts := urlparse(u).path.strip("/").split("/")) > 3
    ]
    if not heads:
        return ""
    head, count = Counter(heads).most_common(1)[0]
    return head if count * 5 >= len(heads) * 4 else ""


def extract(path):
    return extract_sheet(openpyxl.load_workbook(path, data_only=True).worksheets[0])


def extract_sheet(sheet, unnumbered=False):
    """One worksheet's session call and its groups of cases.

    `unnumbered` keeps a case row that has no id, shown as "No id", where the
    default drops it.
    """
    rows = [[clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    header = next(i for i, row in enumerate(rows) if any("Test Case ID" in c for c in row))
    column = columns(rows[header])

    session = ""
    for row in rows[:header]:
        if any(c == "Session API" for c in row):
            session = next((c for c in row if c.startswith("http")), "")

    groups, current, seen = [], None, {}
    for row in rows[header + 1:]:
        def cell(key):
            index = column.get(key)
            return row[index] if index is not None and index < len(row) else ""

        identifier = re.sub(r"\s+", "", cell("id"))
        if not identifier:
            # A group heading: a short number, a name, and who it applies to.
            if len(row) > 1 and row[0] and row[1] and len(row[0]) <= 5:
                applies = " ".join(flat(c) for c in row[2:4] if c)
                # Three sections share one label, so the section number the
                # sheet carries is what tells them apart.
                current = {"label": f"{row[0]}. {flat(row[1])}", "applies": applies, "cases": []}
                groups.append(current)
                continue
            # The M4 sheets head a group with its name alone in the first cell.
            if row and row[0] and not any(row[1:]):
                current = {"label": flat(row[0]), "applies": "", "cases": []}
                groups.append(current)
                continue
            if not (unnumbered and (cell("functionality") or cell("test"))):
                continue
            identifier = "No id"
        if current is None:
            current = {"label": "Test cases", "applies": "", "cases": []}
            groups.append(current)
        anchor = re.sub(r"\s+", "-", identifier.lower())
        seen[anchor] = seen.get(anchor, 0) + 1
        if seen[anchor] > 1:
            anchor = f"{anchor}-{seen[anchor]}"
        current["cases"].append({
            "id": identifier,
            "anchor": anchor,
            "title": flat(cell("functionality")) or identifier,
            "marking": flat(cell("marking")),
            "note": cell("note"),
            "test": cell("test"),
            "steps": cell("steps"),
            "expected": cell("expected"),
            "tester": cell("tester"),
            "apis": list(dict.fromkeys(re.findall(r"https?://[^\s,]+", cell("apis")))),
            "field": flat(cell("field")),
            "applies": flat(cell("applies")),
        })
    for group in groups:
        # A group whose heading names no audience takes it from its rows.
        if not group["applies"]:
            group["applies"] = ", ".join(dict.fromkeys(c["applies"] for c in group["cases"] if c["applies"]))
    return session, [g for g in groups if g["cases"]]


def lines(text):
    return [md(line.strip(" /")) for line in text.split("\n") if line.strip(" /")]


def cell(text):
    """A table cell: one line per line of the sheet, pipes escaped."""
    return "<br/>".join(lines(text)).replace("|", "\\|")


def marking(text, anchors):
    """Shorten the sheet's either-of marking and link the two cases it names."""
    text = MARKINGS.get(text.strip().lower(), re.sub(r"^Yes\b", "Mandatory", text))
    either = re.match(r"Either of the test cases (\S+) or (\S+) is mandatory for Gov\w* Optional for Private", text)
    if not either:
        return cell(text or "Not marked")
    a, b = (f"[{i}](#{anchors.get(i, i.lower())})" for i in either.groups())
    return f"Government: {a} or {b} is mandatory.<br/>Private: optional."


def row(case, routes, anchors, base):
    what = f"**{md(case['title'])}**"
    if case.get("field"):
        what += f"<br/>Field: {md(case['field'])}"
    if case["test"]:
        what += "<br/>" + cell(case["test"])
    if case["steps"]:
        what += f"<details><summary>Steps</summary>{cell(case['steps'])}</details>"
    passes = cell(case["expected"]) or "Not stated."
    if case["tester"]:
        passes += f"<br/>*For the tester:* {cell(case['tester'])}"
    calls = "<br/>".join(call(u, routes, base) for u in case["apis"]) or "None. Checked on your screens or records."
    # Link, not a literal <a>: only Link registers the id with the build's
    # broken anchor check, so a link to a case that is gone fails the build.
    ident = f'<Link id="{case["anchor"]}" to="#{case["anchor"]}">{case["id"]}</Link>'
    return f"| {ident}<br/>{marking(case['marking'], anchors)} | {what} | {passes} | {calls} |"


def render(spec, session, groups, routes):
    total = sum(len(g["cases"]) for g in groups)
    anchors = {c["id"]: c["anchor"] for g in groups for c in g["cases"]}
    base = shared_base(groups)
    out = [
        "---",
        f"title: {spec['title']}",
        f"sidebar_label: {spec['label']}",
        f"sidebar_position: {spec['position']}",
        f"description: The {total} {spec['label']} functional test cases, each with its id, marking, steps, expected result and the APIs it calls.",
        "page_type: reference",
        "toc_max_heading_level: 2",
        "generated: true",
        "---",
        "",
        "import Link from '@docusaurus/Link';",
        "",
        "{/* Generated by scripts/build-test-cases.py from NHA's sheet. Edit the sheet or the script, never this page. */}",
        "",
        f"# {spec['title']}",
        "",
        f"The {total} cases {spec['milestone']} is tested against. They cover {spec['covers']}.",
        "Each row keeps the case id the functional testing report uses, and links every API the case calls.",
        "[How to read a case](/docs/hiecm/v3/resources/test-cases#how-to-read-a-case) explains the columns and markings.",
        "",
    ]
    if session:
        out += [f"Every call needs a gateway session first: {call(session, routes)}.", ""]
    if base:
        out += [f"API paths in the tables below start with `{base}` unless shown in full. The prefix is left off each row to keep the table readable.", ""]
    out += [
        f"An AI coding assistant can walk these cases against your own system with the [{spec['label']} skill]({spec['skill']}).",
        "",
    ]
    level, section = "##", None
    if any(g.get("section") for g in groups):
        level = "###"
    for group in groups:
        if group.get("section") and group["section"] != section:
            section = group["section"]
            out += [f"## {md(section)}", ""]
        out += [f"{level} {md(group['label'])}", ""]
        if group["applies"]:
            out += [f"Applies to: {md(group['applies'])}.", ""]
        # The wrapper's class top-aligns the rows; see .test-cases in mdx.css.
        out += ['<div className="test-cases">', "", "| Case and marking | What is tested | Pass when | APIs |", "| --- | --- | --- | --- |"]
        out += [row(case, routes, anchors, base) for case in group["cases"]]
        out += ["", "</div>", ""]
    out += [
        "## Next steps",
        "",
        "- Certification runs once, for the whole integration: [Go live](/docs/hiecm/v3/getting-started/going-live).",
        f"- The build these cases check: {spec['milestone']}.",
        f"- The next module's cases: {spec['next']}.",
        "",
    ]
    return "\n".join(out)


def main():
    check = "--check" in sys.argv
    routes = load_routes()
    stale = False
    for name, spec in PAGES.items():
        if "sections" in spec:
            session, groups = "", []
            for prefix, path in spec["sections"]:
                book = openpyxl.load_workbook(path, data_only=True)
                for sheet in book.worksheets:
                    title = sheet.title.strip()
                    label = f"{prefix}: {title}" if len(book.worksheets) > 1 else prefix
                    for group in extract_sheet(sheet, unnumbered=True)[1]:
                        group["section"] = label
                        groups.append(group)
        else:
            session, groups = extract(spec["sheet"])
        text = render(spec, session, groups, routes)
        target = OUT / f"{name}.mdx"
        cases = sum(len(g["cases"]) for g in groups)
        if check:
            current = target.read_text() if target.exists() else ""
            if current != text:
                stale = True
            print(f"  {target.relative_to(ROOT)}: {'current' if current == text else 'OUT OF DATE'}")
        else:
            target.write_text(text)
            print(f"  {target.relative_to(ROOT)}: {len(groups)} group(s), {cases} case(s)")
    if stale:
        print("Run: python3 scripts/build-test-cases.py")
        sys.exit(1)


if __name__ == "__main__":
    main()
