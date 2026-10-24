#!/usr/bin/env python3
"""Checks the workbook writer the test skills ship, against NHA's own sheets.

The script integrators copy lives in one place, the python block of
catalogue/shared/concepts/test-workbook-write-back.md, and is folded into the
test skills. This takes that block, not a copy, and for every NHA sheet writes
a result for every case, once as an ft run and once as a ui run, then checks:

- the case ids the script finds are the ones the test case page shows, in the
  same order, including the PHR ids it builds from the rows;
- the four FT columns sit right after NHA's last header (after the added
  Test Case ID column on the PHR sheet);
- every case reads back all four values in its run type's columns: NHA's own
  for ui, or the "(run)" column where NHA had already written in that cell;
- every cell NHA filled still holds exactly what NHA put there, with its
  formatted runs, and every new cell is one the writer was meant to write;
- the merged ranges and the logos are as NHA sent them;
- `open` lists exactly the cases not passed and not marked not run;
- a second write with another tester updates the header block;
- an unknown case id is refused;
- a conditional-pass writes CONDITIONAL PASS, counts as closed in `open`,
  and is refused with empty remarks; a result's steps are appended to Remarks;
- an Aadhaar number, ABHA number, ABHA address, mobile and JWT in a result's
  Response and Remarks reach the sheet masked, and the same result written by
  a copy of the writer with masking taken out fails this check.

Usage: python3 scripts/check-workbook-writeback.py [path to an alternative atom]
Needs openpyxl and Pillow. Exits 0 when every check passes.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText

ROOT = Path(__file__).resolve().parent.parent
ATOM = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "catalogue" / "shared" / "concepts" / "test-workbook-write-back.md"
RAW = ROOT / "catalogue" / "hiecm" / "openapi" / ".raw"
PAGES = ROOT / "site" / "docs" / "hiecm" / "v3" / "resources" / "test-cases"
SHEETS = [
    ("nha-2026-10-22-m1-tests/M1_ABHA_CREATION_VERIFICATION_SCAN_AND_SHARE.xlsx", "m1"),
    ("nha-2026-10-22-m2-tests/M2_BUILDING_HIP_TEST_CASES.xlsx", "m2"),
    ("nha-2026-10-22-m3-tests/M3_BUILDING_HIU_TEST_CASES.xlsx", "m3"),
    ("nha-2026-10-22-m4-tests/HFR_Test_Cases.xlsx", "m4"),
    ("nha-2026-10-22-m4-tests/HPR_Test_Cases.xlsx", "m4"),
    ("nha-2026-10-01-phr-tests/PHR_MOBILE_APP_TEST_CASES.xlsx", "phr"),
]
OUTCOMES = ["passed", "conditional-pass", "failed", "needs-human", "not-run"]
STATUS = {"passed": "PASS", "conditional-pass": "CONDITIONAL PASS", "failed": "FAIL", "needs-human": "", "not-run": "NA"}
# Identity values a result may carry by mistake, and how each must read once masked.
# check-test-report.py carries the same corpus, and checks both scripts mask it alike.
JWT = "eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiJ0ZXN0In0.c2lnbmF0dXJlX3g5"
LEAKS = {"234567890123": "XXXX XXXX 0123", "2345 6789 0126": "XXXX XXXX 0126", "2345-6789-0127": "XXXX XXXX 0127",
         "91-2345-6789-0124": "XX-XXXX-XXXX-0124", "91234567890125": "XX-XXXX-XXXX-0125",
         "asha.verma@sbx": "xxxxxxerma@sbx", "priya_k99@abdm": "xxxxx_k99@abdm", "abc@sbx": "xxx@sbx",
         "9876543210": "XXXXXX3210", "09876543211": "XXXXXX3211", "098765 43212": "XXXXXX3212",
         "987 654 3213": "XXXXXX3213", "98-7654-3214": "XXXXXX3214", "+91 98765 43215": "XXXXXX3215",
         '"firstName": "Asha"': '"firstName": "[masked]"', '"lastName": "Verma"': '"lastName": "[masked]"',
         '"dob": "01-01-1990"': '"dob": "[masked]"', '"yearOfBirth": "1990"': '"yearOfBirth": "[masked]"',
         '"address": "12 MG Road Pune"': '"address": "[masked]"', '"pinCode": "411001"': '"pinCode": "[masked]"',
         '"email": "asha@example.invalid"': '"email": "[masked]"', '"photo": "iVBORw0KGgoAAAANS"': '"photo": "[masked]"',
         "name: Asha Verma": "name: [masked]", '"tokens": {"token": "opaque123x"}': '"tokens": {"token": "[redacted]"}',
         "X-token: Bearer opaque456y": "X-token: [redacted]", "Authorization: Bearer opaque789z": "Authorization: Bearer [redacted]",
         JWT: "[token]"}
# Parts of the corpus that must never appear in any form once masked.
SECRETS = ["opaque123x", "opaque456y", "opaque789z", "Asha", "Verma", "411001", "1990", "iVBORw0", "MG Road",
           "asha@example"]
# Joined with a separator no masking rule reads across.
CORPUS_TEXT = " ; ".join(LEAKS)
UI = [["Status(PASS/FAIL)", "Status"], ["Response of Request with Code"], ["Iterations"], ["Remarks", "Remark"]]
FT = ["FT Status", "FT Response with Code", "FT Iterations", "FT Remarks"]
LABELS = ("Test Executed by:", "Test Execution date:", "Application/ System Name:")
failures = []


def fail(message):
    failures.append(message)
    print(f"FAIL {message}")


def norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def header(ws):
    """Header row, first column of each name, and the last header column."""
    for row in ws.iter_rows(min_row=1, max_row=15):
        names = [norm(c.value) for c in row]
        if "Test Case ID" in names or any(re.fullmatch(r"s\.\s*no", n.lower()) for n in names):
            first = {}
            for name, cell in zip(names, row):
                if name:
                    first.setdefault(name, cell.column)
            return row[0].row, first, max(first.values())
    return None, {}, 0


def merged_anchor(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row, rng.min_col
    return row, col


def cells(wb):
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets
            for row in ws.iter_rows() for c in row if c.value not in (None, "")}


def page_ids(page):
    text = (PAGES / f"{page}.mdx").read_text()
    return [norm(i) for i in re.findall(r'<Link id="[^"]+" to="#[^"]+">([^<]+)</Link>', text)]


def main():
    block = re.search(r"```python\n(.*?)\n```", ATOM.read_text(), re.S)
    if not block:
        sys.exit(f"no python block in {ATOM.relative_to(ROOT)}")
    work = Path(tempfile.mkdtemp())
    script = work / "abdm_workbook.py"
    script.write_text(block.group(1) + "\n")

    def run(*args):
        return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True)

    m4_ids = []
    for rel, page in SHEETS:
        src = RAW / rel
        book = work / src.name
        shutil.copy2(src, book)
        fails_before = len(failures)
        listed = run("cases", book)
        if listed.returncode:
            fail(f"{rel}: cases exited {listed.returncode}: {listed.stderr.strip()}")
            continue
        ids = [i for sheet in json.loads(listed.stdout).values() for i in sheet]
        shown = [re.sub(r"#\d+$", "", i).replace(" ", "") for i in ids]
        if page == "m4":
            m4_ids += shown
        else:
            want = [i.replace(" ", "") for i in page_ids(page)]
            if shown != want:
                fail(f"{rel}: the script finds {len(ids)} ids, the page shows {len(want)}, or not in the same order")

        nha = openpyxl.load_workbook(src, rich_text=True)
        results = {r: [{"id": i, "outcome": OUTCOMES[n % 5], "response": f"{r} response {n}",
                        "iterations": n % 3 + 1, "remarks": f"{r} remark {n}"} for n, i in enumerate(ids)]
                   for r in ("ft", "ui")}
        for r in ("ft", "ui"):
            path = work / f"{r}.json"
            path.write_text(json.dumps(results[r]))
            out = run("write", book, "--run", r, "--results", path, "--executed-by", "Tester A", "--application", "App")
            if out.returncode or out.stdout or out.stderr:
                fail(f"{rel} {r}: write exited {out.returncode}: {out.stderr.strip()}")
        if not (work / f"{src.stem}.before-abdm-run{src.suffix}").exists():
            fail(f"{rel}: no backup written")

        got = openpyxl.load_workbook(book, rich_text=True)
        before, after = cells(nha), cells(got)
        for key, value in before.items():
            if norm(value).startswith(LABELS):
                continue
            if after.get(key) != value:
                fail(f"{rel}: {key[0]}!{key[1]} NHA's value changed")
                break
        rich_before = sum(isinstance(v, CellRichText) for v in before.values())
        rich_after = sum(isinstance(v, CellRichText) for v in after.values())
        if rich_after < rich_before:
            fail(f"{rel}: {rich_before - rich_after} formatted cell(s) lost their runs")
        for ws in nha.worksheets:
            a, b = ws, got[ws.title]
            if sorted(map(str, a.merged_cells.ranges)) != sorted(map(str, b.merged_cells.ranges)) \
                    or len(a._images) != len(b._images):
                fail(f"{rel}: {ws.title} merged ranges or logos changed")

        # Where every value should be, from the written sheet's own header.
        expected = {}
        by_id = {}
        for ws in got.worksheets:
            top, cols, _ = header(ws)
            if top is None or "Test Case ID" not in cols:
                continue
            _, nha_cols, nha_last = header(nha[ws.title])
            start = nha_last + 1 if "Test Case ID" in nha_cols else nha_last + 2
            for k, name in enumerate(FT):
                if cols.get(name) != start + k:
                    fail(f"{rel}: {ws.title} '{name}' is in column {cols.get(name)}, not {start + k}")
            if "Test Case ID" not in nha_cols:
                expected[(ws.title, ws.cell(top, cols["Test Case ID"]).coordinate)] = "Test Case ID"
            for name in [n for names in UI for n in names] + FT + [n for n in cols if n.endswith(" (run)")]:
                if name in cols:
                    expected.setdefault((ws.title, ws.cell(top, cols[name]).coordinate), name)
            n = {}
            for r in range(top + 1, ws.max_row + 1):
                value = norm(ws.cell(r, cols["Test Case ID"]).value).replace(" ", "")
                if not value:
                    continue
                n[value] = n.get(value, 0) + 1
                key = value if n[value] == 1 else f"{value}#{n[value]}"
                by_id[key] = (ws, top, cols, r)
                if "Test Case ID" not in nha_cols:
                    expected[(ws.title, ws.cell(r, cols["Test Case ID"]).coordinate)] = value
        for r, res_list in results.items():
            for res in res_list:
                if res["id"] not in by_id:
                    fail(f"{rel}: {res['id']} has no row")
                    continue
                ws, top, cols, row = by_id[res["id"]]
                remark = ("Needs tester: " + res["remarks"]) if res["outcome"] == "needs-human" else res["remarks"]
                values = [STATUS[res["outcome"]], res["response"], res["iterations"], remark]
                for k in range(4):
                    if r == "ft":
                        col = cols[FT[k]]
                    else:
                        col = next(cols[n] for n in UI[k] if n in cols)
                        arow, acol = merged_anchor(ws, row, col)
                        if norm(nha[ws.title].cell(arow, acol).value):
                            col = cols.get(f"{norm(ws.cell(top, col).value)} (run)")
                            if col is None:
                                fail(f"{rel}: {ws.title} row {row} holds NHA text but has no '(run)' column")
                                continue
                    arow, acol = merged_anchor(ws, row, col)
                    coord = ws.cell(arow, acol).coordinate
                    expected[(ws.title, coord)] = values[k]
        for key, value in expected.items():
            if key in before:
                continue
            if norm(after.get(key)) != norm(value):
                fail(f"{rel}: {key[0]}!{key[1]} reads {after.get(key)!r}, expected {value!r}")
        for key, value in after.items():
            if key in before or key in expected:
                continue
            fail(f"{rel}: {key[0]}!{key[1]} = {value!r} written outside its place")

        for r in ("ft", "ui"):
            out = run("open", book, "--run", r)
            want_open = {x["id"] for x in results[r] if x["outcome"] not in ("passed", "conditional-pass", "not-run")}
            if set(out.stdout.split()) != want_open:
                fail(f"{rel} {r}: open listed {len(set(out.stdout.split()))} cases, expected {len(want_open)}")

        if any(norm(v).startswith(LABELS[0]) for v in before.values()):
            again = run("write", book, "--run", "ft", "--results", work / "ft.json", "--executed-by", "Tester B")
            labels = [norm(v) for v in cells(openpyxl.load_workbook(book)).values() if norm(v).startswith(LABELS[0])]
            if again.returncode or labels != [f"{LABELS[0]} Tester B"] * len(labels):
                fail(f"{rel}: a second write did not update the tester: {labels}")
        state = "ok  " if len(failures) == fails_before else "FAIL"
        print(f"{state} {rel}: {len(ids)} case(s)")

    want_m4 = [i.replace(" ", "") for i in page_ids("m4") if i != "No id"]
    if m4_ids != want_m4:
        fail(f"M4: the script finds {len(m4_ids)} ids across HFR and HPR, the page shows {len(want_m4)}")

    # A workbook where someone had already typed PASS in NHA's own Status: a
    # needs-human result goes in the "(run)" column and must stay open.
    pre = work / "prefilled.xlsx"
    shutil.copy2(RAW / SHEETS[0][0], pre)
    wb = openpyxl.load_workbook(pre)
    ws = wb.worksheets[0]
    top, cols, _ = header(ws)
    row = next(r for r in range(top + 1, ws.max_row + 1) if norm(ws.cell(r, cols["Test Case ID"]).value) == "CRT_ABHA_101")
    ws.cell(row, cols["Status(PASS/FAIL)"]).value = "PASS"
    wb.save(pre)
    (work / "pre.json").write_text(json.dumps([{"id": "CRT_ABHA_101", "outcome": "needs-human", "remarks": "OTP"}]))
    run("write", pre, "--run", "ui", "--results", work / "pre.json")
    if "CRT_ABHA_101" not in run("open", pre, "--run", "ui").stdout.split():
        fail("a needs-human case beside NHA's own PASS is not listed as open")
    else:
        print("ok   a needs-human case beside NHA's own PASS stays open")

    # Conditional pass, steps and masking, on a fresh M1 copy, read back from the ui columns.
    def ui_cells(book, case):
        wb = openpyxl.load_workbook(book)
        nha = openpyxl.load_workbook(RAW / SHEETS[0][0])
        ws = wb.worksheets[0]
        top, cols, _ = header(ws)
        row = next(r for r in range(top + 1, ws.max_row + 1) if norm(ws.cell(r, cols["Test Case ID"]).value).replace(" ", "") == case)
        out = {}
        for names in UI:
            col = next(cols[n] for n in names if n in cols)
            if norm(nha[ws.title].cell(*merged_anchor(ws, row, col)).value):
                col = cols[f"{norm(ws.cell(top, col).value)} (run)"]
            out[names[0]] = norm(ws.cell(*merged_anchor(ws, row, col)).value)
        return out

    def privacy(script_path):
        """What a writer leaves unmasked in the sheet, for one result carrying every kind of identity."""
        book = work / "privacy.xlsx"
        shutil.copy2(RAW / SHEETS[0][0], book)
        for old in work.glob("privacy.before-abdm-run*"):
            old.unlink()
        text = CORPUS_TEXT
        (work / "privacy.json").write_text(json.dumps([{"id": "CRT_ABHA_112", "outcome": "failed", "iterations": 1,
                                                       "response": "HTTP 200 " + text, "remarks": "returned " + text}]))
        out = subprocess.run([sys.executable, str(script_path), "write", book, "--run", "ui", "--results",
                              work / "privacy.json"], capture_output=True, text=True)
        if out.returncode:
            return ["write exited " + str(out.returncode) + ": " + out.stderr.strip()]
        got = ui_cells(book, "CRT_ABHA_112")
        found = []
        for name in ("Response of Request with Code", "Remarks"):
            for raw, masked in LEAKS.items():
                if raw in got[name] or masked not in got[name]:
                    found.append(f"{name} carries {raw!r} unmasked")
            found += [f"{name} carries {x!r}" for x in SECRETS if x in got[name]]
        return found

    leaks = privacy(script)
    if leaks:
        for leak in leaks:
            fail("masking: " + leak)
    else:
        print("ok   every value in the identity corpus reaches the sheet masked")
    broken = work / "broken_workbook.py"
    source = script.read_text()
    broken.write_text(source.replace('\nif __name__ == "__main__":', '\nmask = str\n\nif __name__ == "__main__":', 1))
    if broken.read_text() == source or not privacy(broken):
        fail("masking: a writer with masking taken out is not caught")
    else:
        print("ok   a writer with masking taken out fails the masking check")

    book = work / "conditional.xlsx"
    shutil.copy2(RAW / SHEETS[0][0], book)
    (work / "cond.json").write_text(json.dumps([
        {"id": "CRT_ABHA_113", "outcome": "conditional-pass", "iterations": 1, "remarks": "Shown; the card has no QR"},
        {"id": "SHARE_PATIENT_PROFILE_702", "outcome": "needs-human", "iterations": 1, "remarks": "a second ABHA address",
         "steps": [{"step": 1, "outcome": "passed"}, {"step": 2, "outcome": "needs-human", "note": "a second ABHA address"},
                   {"step": 3, "outcome": "passed"}]}]))
    out = run("write", book, "--run", "ui", "--results", work / "cond.json")
    cond, steps = ui_cells(book, "CRT_ABHA_113"), ui_cells(book, "SHARE_PATIENT_PROFILE_702")
    still = run("open", book, "--run", "ui").stdout.split()
    if out.returncode or cond["Status(PASS/FAIL)"] != "CONDITIONAL PASS" or cond["Remarks"] != "Shown; the card has no QR" \
            or "CRT_ABHA_113" in still or "SHARE_PATIENT_PROFILE_702" not in still:
        fail(f"conditional pass: written {cond}, open lists it: {'CRT_ABHA_113' in still}")
    else:
        print("ok   a conditional-pass writes CONDITIONAL PASS with its condition, and open counts it closed")
    want = "Needs tester: a second ABHA address Steps: 1 PASS; 2 needs tester (a second ABHA address); 3 PASS"
    if steps["Remarks"] != want:
        fail(f"steps: Remarks reads {steps['Remarks']!r}, expected {want!r}")
    else:
        print("ok   a result's steps are appended to Remarks")
    (work / "nocond.json").write_text(json.dumps([{"id": "CRT_ABHA_113", "outcome": "conditional-pass", "remarks": " "}]))
    out = run("write", book, "--run", "ui", "--results", work / "nocond.json")
    if out.returncode == 0 or "conditional-pass needs remarks" not in out.stderr:
        fail("a conditional-pass with empty remarks is not refused")
    else:
        print("ok   a conditional-pass with empty remarks is refused")

    # Steps that disagree with the case's outcome are refused.
    bad_steps = [("passed", ["passed", "failed"]), ("passed", ["passed", "needs-human"]),
                 ("conditional-pass", ["passed", "failed"]), ("needs-human", ["passed", "failed", "needs-human"]),
                 ("needs-human", ["passed", "not-run"]), ("not-run", ["failed"])]
    refused = []
    for outcome, seen in bad_steps:
        (work / "badsteps.json").write_text(json.dumps([{"id": "CRT_ABHA_113", "outcome": outcome, "remarks": "x",
                                                         "steps": [{"step": n + 1, "outcome": o} for n, o in enumerate(seen)]}]))
        out = run("write", book, "--run", "ui", "--results", work / "badsteps.json")
        refused.append(out.returncode != 0 and "steps disagree" in out.stderr)
    good = [("failed", ["passed", "failed", "not-run"]), ("needs-human", ["passed", "needs-human", "not-run"]),
            ("conditional-pass", ["passed", "needs-human"]), ("passed", ["passed", "passed"])]
    accepted = []
    for outcome, seen in good:
        (work / "goodsteps.json").write_text(json.dumps([{"id": "CRT_ABHA_113", "outcome": outcome, "remarks": "x",
                                                          "steps": [{"step": n + 1, "outcome": o} for n, o in enumerate(seen)]}]))
        accepted.append(run("write", book, "--run", "ui", "--results", work / "goodsteps.json").returncode == 0)
    if all(refused) and all(accepted):
        print("ok   steps that disagree with the outcome are refused, steps that agree are written")
    else:
        fail(f"steps against the outcome: refused {refused}, accepted {accepted}")

    # open: statuses in any case and spacing, and with --entity an NA on a mandatory case.
    book = work / "open.xlsx"
    shutil.copy2(RAW / SHEETS[0][0], book)
    every = json.loads(run("cases", book).stdout)
    ids = [i for sheet in every.values() for i in sheet]
    res = [{"id": i, "outcome": "passed"} for i in ids]
    for r in res:
        if r["id"] in ("CRT_ABHA_101", "CRT_ABHA_301", "CRT_ABHA_103"):
            r["outcome"] = "not-run"
    (work / "open.json").write_text(json.dumps(res))
    run("write", book, "--run", "ui", "--results", work / "open.json")
    wb = openpyxl.load_workbook(book)
    ws = wb.worksheets[0]
    top, cols, _ = header(ws)
    for r in range(top + 1, ws.max_row + 1):
        cid = norm(ws.cell(r, cols["Test Case ID"]).value).replace(" ", "")
        if cid == "CRT_ABHA_104":
            ws.cell(r, cols["Status(PASS/FAIL)"]).value = " pass "
        if cid == "CRT_ABHA_105":
            ws.cell(r, cols["Status(PASS/FAIL)"]).value = "Conditional Pass"
    wb.save(book)

    def listed(*extra):
        return {line.split()[0]: line for line in run("open", book, "--run", "ui", *extra).stdout.splitlines() if line}

    plain, priv, gov = listed(), listed("--entity", "private"), listed("--entity", "government")
    if plain or set(priv) != {"CRT_ABHA_101"} or "NA on a mandatory case" not in priv.get("CRT_ABHA_101", "") \
            or set(gov) != {"CRT_ABHA_101", "CRT_ABHA_301"}:
        fail(f"open: without --entity {sorted(plain)}, private {sorted(priv)}, government {sorted(gov)}")
    else:
        print("ok   open reads ' pass ' and 'Conditional Pass' as closed, and with --entity lists an NA on a mandatory case")

    (work / "missing.json").write_text(json.dumps([{"id": "NOT_A_CASE", "outcome": "passed"}]))
    bad = run("write", work / Path(SHEETS[2][0]).name, "--run", "ft", "--results", work / "missing.json")
    if bad.returncode != 1 or "not in the workbook: NOT_A_CASE" not in bad.stderr:
        fail("an unknown case id is not refused")
    else:
        print("ok   an unknown case id is refused")

    shutil.rmtree(work)
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
