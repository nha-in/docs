#!/usr/bin/env python3
"""Checks the tracker and the findings file the M1 skill writes, on synthetic results.

The report script lives in the python block of
catalogue/shared/concepts/test-report.md; functionalities.json is the file
scripts/build-skills.mjs ships beside it in the M1 skill. This takes the
block, not a copy, fills a copy of NHA's M1 workbook with made up results
through the writer in catalogue/shared/concepts/test-workbook-write-back.md,
and checks:

- functionalities.json agrees with the M1 skill: the same functionalities in
  the same order, each with the cases its file's Cases table lists, and each
  cited case sits in a table row of that file's Screen requirements;
- tracker refuses to run with no entity type, and refuses an entity type
  that disagrees with the one it holds;
- it writes INTEGRATION-PLAN.md with one row per functionality between the
  two markers, keeps what the plan holds outside them, and replaces the block
  on a second run rather than adding one;
- each functionality's state follows the workbook: passed, failing,
  built and waiting for the tester, not started, optional and not
  applicable; --set building holds over a FAIL until --set tested, and
  --screen lands in the Screen or route column;
- its failing and open mandatory cases are the ones check lists;
- findings lists every FAIL and CONDITIONAL PASS once, with its functionality,
  says yes where a screen requirement cites the case, and writes a Catalogue
  issue ready to paste where none does, naming the case and the functionality;
- build writes the findings beside the draft when functionalities.json sits
  beside the script;
- no identity value from a result reaches the plan, the tracker state or the
  findings.

Usage: python3 scripts/check-tracker-findings.py [path to an alternative report atom]
Needs openpyxl, Pillow, python-docx and reportlab. Exits 0 when every check passes.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
REPORT_ATOM = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "catalogue" / "shared" / "concepts" / "test-report.md"
WRITER_ATOM = ROOT / "catalogue" / "shared" / "concepts" / "test-workbook-write-back.md"
SKILL = ROOT / "plugins" / "abdm-integrators-assistant" / "skills" / "abdm-m1"
M1 = ROOT / "catalogue" / "hiecm" / "openapi" / ".raw" / "nha-2026-10-22-m1-tests" / "M1_ABHA_CREATION_VERIFICATION_SCAN_AND_SHARE.xlsx"
# Made up identity values a result may carry by mistake. None may reach a tracker or findings file.
LEAKS = ["234567890123", "91-2345-6789-0124", "asha.verma@sbx", "9876543210", "Asha Verma"]
failures = []


def ok(cond, message):
    if cond:
        print(f"ok   {message}")
    else:
        failures.append(message)
        print(f"FAIL {message}")


def block(path):
    found = re.search(r"```python\n(.*?)\n```", path.read_text(), re.S)
    if not found:
        sys.exit(f"no python block in {path}")
    return found.group(1) + "\n"


def section(markdown, heading):
    parts = re.split(r"\n(?=## )", markdown)
    return next((p for p in parts if p.startswith("## " + heading + "\n")), "")


def main():
    data = json.loads((SKILL / "scripts" / "functionalities.json").read_text())
    router = (SKILL / "SKILL.md").read_text()
    order = re.findall(r"\(references/functionalities/([a-z0-9-]+)\.md\)", section(router, "Functionalities, in build order"))
    ok([f["id"] for f in data["functionalities"]] == order,
       "functionalities.json: the functionalities of the router's table, in its order")
    agree = []
    for f in data["functionalities"]:
        text = (SKILL / f["file"]).read_text()
        listed = re.findall(r"^\| ([A-Z][A-Z0-9_]+) \|", section(text, "Cases"), re.M)
        rows = [line for line in section(text, "Screen requirements").splitlines() if line.startswith("|")]
        if listed != f["cases"]:
            agree.append(f["id"] + " cases")
        for req in f["screen_requirements"]:
            if not any(all(c in row for c in req["cases"]) and req["requirement"][:40] in row.replace("\\|", "|") for row in rows):
                agree.append(f["id"] + " requirement")
    ok(not agree, f"functionalities.json: each file's cases, and each requirement with its cases in that file: {agree[:3]}")

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        ft = work / "abdm-ft"
        ft.mkdir()
        report = ft / "abdm_report.py"
        report.write_text(block(REPORT_ATOM))
        writer = ft / "abdm_workbook.py"
        writer.write_text(block(WRITER_ATOM))
        book = ft / M1.name
        shutil.copy2(M1, book)
        plan = work / "INTEGRATION-PLAN.md"
        plan.write_text("# Integration plan\n\nGoal: M1 for a private hospital. Hand written line kept.\n")
        leak = " ".join(LEAKS)

        def run(*args, script=report):
            return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True, cwd=work)

        def write(results, target=book):
            res = work / "results.json"
            res.write_text(json.dumps(results))
            out = run("write", target, "--run", "ui", "--results", res, script=writer)
            if out.returncode:
                sys.exit("the writer failed: " + out.stderr.strip())

        def states_of(out):
            return dict(line.split(": ", 1) for line in out.stdout.splitlines() if re.match(r"^[a-z0-9-]+: ", line))

        def agrees(out, entity, *meta):
            check = run("check", book, "--run", "ui", "--entity", entity, *meta)
            blocking = len(re.findall(r"^not passed ", check.stdout, re.M))
            m = re.search(r"mandatory cases (\d+), (\d+) passed, (\d+) failing, (\d+) open", out.stdout)
            return bool(m) and int(m.group(3)) + int(m.group(4)) == blocking, blocking

        write([
            {"id": "CRT_ABHA_101", "outcome": "passed", "response": "HTTP 200", "remarks": "option shown"},
            {"id": "CRT_ABHA_104", "outcome": "failed", "response": "HTTP 400 " + leak, "remarks": "accepted 11 digits " + leak},
            {"id": "CRT_ABHA_103", "outcome": "failed", "response": "HTTP 200", "remarks": "one language only"},
            {"id": "CRT_ABHA_206", "outcome": "failed", "response": "HTTP 200", "remarks": "mobile step skipped"},
            {"id": "VRFY_ABHA_101", "outcome": "conditional-pass", "response": "HTTP 200", "remarks": "resend timer not shown"},
            {"id": "VRFY_ABHA_202", "outcome": "passed", "response": "HTTP 200", "remarks": "to be emptied"},
            {"id": "VRFY_ABHA_301", "outcome": "needs-human", "response": "", "remarks": "a mobile with no ABHA"},
            {"id": "TAGGING_UNIQUEPATIENTID_UNIQUEABHANUMBER", "outcome": "not-run", "response": "", "remarks": "not offered"},
            {"id": "CRT_ABHA_115", "outcome": "passed", "response": "HTTP 200", "remarks": "own card"},
        ])
        # A CONDITIONAL PASS with no condition, typed into the sheet by hand, as the gate's own check does.
        wb = openpyxl.load_workbook(book)
        ws = wb.worksheets[0]
        head = {str(c.value): c.column for c in ws[9] if c.value}
        for r in range(10, ws.max_row + 1):
            if str(ws.cell(r, head["Test Case ID"]).value or "").strip() == "VRFY_ABHA_202":
                ws.cell(r, head["Status(PASS/FAIL)"]).value = "CONDITIONAL PASS"
                ws.cell(r, head["Remarks"]).value = None
        wb.save(book)
        meta = work / "agreed.json"
        meta.write_text(json.dumps({"agreed_not_run": {"TAGGING_UNIQUEPATIENTID_UNIQUEABHANUMBER": "<NAME>, NHA, email of 2026-10-24"}}))

        # The tracker, with no functionalities.json beside the script.
        out = run("tracker", book, "--run", "ui", "--entity", "private", "--plan", plan)
        ok(out.returncode != 0 and "not built per functionality" in (out.stderr + out.stdout),
           "tracker: refuses without functionalities.json, naming why")
        shutil.copy2(SKILL / "scripts" / "functionalities.json", ft / "functionalities.json")
        out = run("tracker", book, "--run", "ui", "--plan", plan)
        ok(out.returncode != 0 and "no entity type" in (out.stderr + out.stdout), "tracker: refuses with no entity type, asking for it")
        out = run("tracker", book, "--run", "ui", "--entity", "private", "--plan", plan)
        states = states_of(out)
        text = plan.read_text()
        rows = [line for line in text.splitlines() if re.match(r"^\| \d+ \|", line)]
        ok(out.returncode == 0 and len(rows) == len(data["functionalities"]) and text.count("<!-- abdm-tracker start") == 1
           and "Hand written line kept." in text, "tracker: one row per functionality between the markers, the rest of the plan kept")
        expect = {"create-abha-aadhaar-otp": "failing (CRT_ABHA_104 FAIL)", "login-abha-number": "built and waiting for the tester",
                  "login-abha-address": "built and waiting for the tester", "find-abha-mobile": "built and waiting for the tester",
                  "link-one-abha-per-patient": "built and waiting for the tester", "login-aadhaar-otp": "not started",
                  "update-mobile": "optional, not in scope", "create-abha-demo-auth": "not applicable"}
        wrong = {k: states.get(k) for k, v in expect.items() if states.get(k) != v}
        ok(not wrong, f"tracker: each state follows the workbook, and an NA nobody agreed does not pass: {wrong}")
        same, blocking = agrees(out, "private")
        ok(same, f"tracker: its failing and open mandatory cases are the {blocking} check lists")
        ok("VRFY_ABHA_202 CONDITIONAL PASS with no condition" in text and "CRT_ABHA_104 FAIL" in text
           and "VRFY_ABHA_101 CONDITIONAL PASS" in text and "VRFY_ABHA_301 needs tester" in text
           and "TAGGING_UNIQUEPATIENTID_UNIQUEABHANUMBER NA" in text,
           "tracker: each mandatory case shows its result, a CONDITIONAL PASS with no condition as not through")

        out = run("tracker", book, "--run", "ui", "--plan", plan, "--meta", meta)
        same, blocking = agrees(out, "private", "--meta", meta)
        ok(states_of(out).get("link-one-abha-per-patient") == "passed" and "NA, agreed" in plan.read_text() and same,
           f"tracker: an NA NHA agreed closes the case, as the gate does, and both list the same {blocking} open")

        # Held states, and when they give way.
        out = run("tracker", book, "--run", "ui", "--plan", plan, "--set", "create-abha-aadhaar-otp=building",
                  "--screen", "CRT_ABHA_104=/patients/new, Aadhaar field")
        text = plan.read_text()
        ok("create-abha-aadhaar-otp: building" in out.stdout and "CRT_ABHA_104: /patients/new, Aadhaar field" in text
           and text.count("<!-- abdm-tracker start") == 1,
           "tracker: --set building holds over a FAIL, --screen lands in its column, and the block is replaced, not added")
        out = run("tracker", book, "--run", "ui", "--plan", plan)
        ok("create-abha-aadhaar-otp: building" in out.stdout, "tracker: a held state holds while the workbook is unchanged")
        write([{"id": "CRT_ABHA_107", "outcome": "failed", "response": "HTTP 422", "remarks": "wrong OTP accepted"}])
        out = run("tracker", book, "--run", "ui", "--plan", plan)
        ok("create-abha-aadhaar-otp: failing" in out.stdout,
           "tracker: a held state gives way when a later write leaves the functionality failing")
        out = run("tracker", book, "--run", "ui", "--plan", plan, "--set", "create-abha-aadhaar-otp=tested",
                  "--build", "build 42 for 9876543210")
        text = plan.read_text()
        ok("create-abha-aadhaar-otp: failing" in out.stdout and "build 42 for XXXXXX3210" in text and "9876543210" not in text,
           "tracker: --set tested hands the state to the workbook and records the build, masked")
        out = run("tracker", book, "--run", "ui", "--entity", "government", "--plan", plan)
        ok(out.returncode != 0 and "ask the tester" in (out.stderr + out.stdout),
           "tracker: refuses an entity type that disagrees with the one it holds")

        # A government entity: CRT_ABHA_114 or CRT_ABHA_115 is one either pair.
        gov_plan = work / "GOV-PLAN.md"
        out = run("tracker", book, "--run", "ui", "--entity", "government", "--plan", gov_plan, "--state", work / "gov.json")
        same, blocking = agrees(out, "government")
        ok("CRT_ABHA_114 open or CRT_ABHA_115 PASS" in gov_plan.read_text() and same,
           f"tracker: a government entity's either pair is one case, met by CRT_ABHA_115, as check counts it ({blocking} open)")

        # The findings, with the shipped file.
        out = run("findings", book, "--run", "ui", "--entity", "private", "--out", ft / "draft-report-m1.pdf")
        findings = (ft / "draft-report-m1-findings.md").read_text()
        listed = re.findall(r"^\| ([A-Z][A-Z0-9_]+) \| (FAIL|CONDITIONAL PASS) \|", findings, re.M)
        ok(out.returncode == 0 and sorted(listed) == sorted([("CRT_ABHA_103", "FAIL"), ("CRT_ABHA_104", "FAIL"),
                                                             ("CRT_ABHA_107", "FAIL"), ("CRT_ABHA_206", "FAIL"),
                                                             ("VRFY_ABHA_101", "CONDITIONAL PASS"),
                                                             ("VRFY_ABHA_202", "CONDITIONAL PASS")]),
           "findings: every FAIL and CONDITIONAL PASS once, optional ones included, and nothing else")
        row206 = next((line for line in findings.splitlines() if line.startswith("| CRT_ABHA_206 ")), "")
        ok("yes, through create-abha-aadhaar-otp, its rule for CRT_ABHA_108" in row206,
           "findings: CRT_ABHA_206, handed to create-abha-aadhaar-otp, is covered only through its rule for CRT_ABHA_108, named")
        ok("`create-abha-aadhaar-otp`" in findings and "| yes: " in findings and "None: every FAIL" in findings,
           "findings: each names its functionality and the rule stating its check; no issue when every case has one")

        # A file whose rows no longer state a check for CRT_ABHA_104 or CRT_ABHA_108.
        bare = json.loads(json.dumps(data))
        for f in bare["functionalities"]:
            for req in f["screen_requirements"]:
                if not req.get("delegates_to"):
                    req["cases"] = [c for c in req["cases"] if c not in ("CRT_ABHA_104", "CRT_ABHA_108")]
        bare_path = work / "bare.json"
        bare_path.write_text(json.dumps(bare))
        out = run("findings", book, "--run", "ui", "--entity", "private", "--out", ft / "bare-report.pdf",
                  "--functionalities", bare_path)
        findings = (ft / "bare-report-findings.md").read_text()
        issue = findings.split("## Catalogue issues to file", 1)[-1]
        ok("catalogue issue CRT_ABHA_104 in create-abha-aadhaar-otp" in out.stdout
           and "### CRT_ABHA_104: no screen requirement in create-abha-aadhaar-otp" in issue
           and "Title: create-abha-aadhaar-otp: no screen requirement states the check for CRT_ABHA_104" in issue
           and "references/functionalities/create-abha-aadhaar-otp.md" in issue and "Not filed." in issue
           and "CRT_ABHA_103:" not in issue,
           "findings: a case with no rule gets a Catalogue issue ready to paste, naming case, functionality and file")
        ok("catalogue issue CRT_ABHA_206 in create-abha-biometric" in out.stdout
           and "Why: it hands the case to create-abha-aadhaar-otp" in issue,
           "findings: a delegation row alone does not cover a case: CRT_ABHA_206 is an issue once its target states no check")

        # build writes the findings beside the draft.
        meta2 = work / "meta.json"
        meta2.write_text(json.dumps({"organisation": "<ORG>", "product": "<PRODUCT>", "milestone": "M1", "entity_type": "private",
                                     "role": "HIP", "executed_by": "<TESTER>", "findings": {}}))
        out = run("build", book, "--run", "ui", "--meta", meta2, "--evidence", ft / "evidence", "--out", ft / "draft.pdf",
                  "--draft", "--format", "docx")
        ok(out.returncode == 0 and (ft / "draft-findings.md").exists() and "findings: 6 FAIL or CONDITIONAL PASS" in out.stdout,
           "build: writes <report>-findings.md beside the draft")

        written = plan.read_text() + gov_plan.read_text() + (ft / "tracker.json").read_text() + \
            "".join(p.read_text() for p in ft.glob("*-findings.md"))
        leaked = [x for x in LEAKS if x in written]
        ok(not leaked, f"no identity value from a result reaches the plan, the tracker state or the findings: {leaked}")

    if failures:
        print(f"\n{len(failures)} check(s) failed")
        sys.exit(1)
    print("\nThe tracker and the findings pass every check.")


if __name__ == "__main__":
    main()
